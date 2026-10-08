"""CSV/XLSX ingestion with bounded parsing and original-row preservation."""

import csv
import io
import json
import zipfile
from collections import OrderedDict
from decimal import Decimal
from pathlib import Path

import openpyxl

from .models import Invoice, LineItem, Party, Table
from .normalize import HEADER_MAP, header, iso_date, number, text
from .validation import TAX_KEYS, ZERO

MAX_ROWS = 5000
MAX_COLUMNS = 100


def read_sheets(path: Path, kind: str) -> list[tuple[str, list[list]]]:
    if kind == "csv":
        data = path.read_bytes()
        try:
            contents = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                contents = data.decode("utf-16")
            except UnicodeDecodeError as exc:
                raise ValueError("CSV must use UTF-8 or UTF-16 encoding.") from exc
        try:
            delimiter = csv.Sniffer().sniff(contents[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            # Export preambles often defeat frequency-based sniffing. Prefer the
            # delimiter that reveals actual known column names, not a guessed comma.
            def header_score(candidate):
                try:
                    preview = list(csv.reader(io.StringIO(contents[:8192]), delimiter=candidate))[
                        :30
                    ]
                    return max(
                        (sum(header(cell) in HEADER_MAP for cell in row) for row in preview),
                        default=0,
                    )
                except csv.Error:
                    return 0

            delimiter = max(",;\t|", key=header_score)
        rows = []
        try:
            for row in csv.reader(io.StringIO(contents), delimiter=delimiter):
                if len(row) > MAX_COLUMNS:
                    raise ValueError("Maximum 100 columns supported.")
                rows.append(row)
                if len(rows) > MAX_ROWS:
                    raise ValueError("Maximum 5,000 rows per document supported.")
        except csv.Error as exc:
            raise ValueError(f"Malformed CSV: {exc}") from exc
        return [("CSV", rows)]
    try:
        with zipfile.ZipFile(path) as archive:
            if (
                len(archive.infolist()) > 2000
                or sum(e.file_size for e in archive.infolist()) > 80_000_000
            ):
                raise ValueError("Workbook expanded size exceeds the safe processing limit.")
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError("Not an XLSX workbook.")
        workbook = openpyxl.load_workbook(
            io.BytesIO(path.read_bytes()), read_only=True, data_only=False, keep_links=False
        )
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValueError("Invalid XLSX workbook.") from exc
    sheets = []
    count = 0
    try:
        if len(workbook.worksheets) > 20:
            raise ValueError("Maximum 20 worksheets supported.")
        for sheet in workbook.worksheets:
            if (sheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError("Maximum 100 columns supported.")
            rows = []
            for row in sheet.iter_rows(values_only=True):
                count += 1
                if count > MAX_ROWS:
                    raise ValueError("Maximum 5,000 rows per document supported.")
                rows.append(list(row))
            sheets.append((sheet.title, rows))
    finally:
        workbook.close()
    return sheets


def extract_tabular(path: Path, kind: str):
    invoices, tables, warnings = [], [], []
    for sheet, rows in read_sheets(path, kind):
        rows = [row for row in rows if any(text(cell) is not None for cell in row)]
        if not rows:
            continue
        scores = [len({HEADER_MAP.get(header(cell)) for cell in row} - {None}) for row in rows[:30]]
        header_index = max(range(len(scores)), key=scores.__getitem__)
        if scores[header_index] < 2:
            raise ValueError(
                f"{sheet}: could not recognize a table header. Include fields such as Invoice No, Description, Taxable Value, GST and Grand Total."
            )
        if header_index:
            warnings.append(
                f"{sheet}: {header_index} pre-header rows excluded from invoice mapping."
            )
        names, mappings = [], []
        width = max(len(row) for row in rows[header_index:])
        header_row = rows[header_index] + [None] * (width - len(rows[header_index]))
        for n, value in enumerate(header_row):
            name = text(value) or f"column_{n + 1}"
            if name in names:
                name = f"{name}_{n + 1}"
            names.append(name)
            mappings.append(HEADER_MAP.get(header(value)))
        known = [key for key in mappings if key]
        if len(known) != len(set(known)):
            warnings.append(
                f"{sheet}: multiple columns map to one field; the last nonempty value is used. Review the preserved table."
            )
        original_rows = []
        groups = OrderedDict()
        current_key = None
        for row_index, row in enumerate(rows[header_index + 1 :], header_index + 2):
            clean = {name: text(row[n]) if n < len(row) else None for n, name in enumerate(names)}
            original_rows.append(clean)
            mapped = {}
            for name, key in zip(names, mappings):
                value = clean[name]
                if value and value.startswith("="):
                    warnings.append(
                        f"{sheet} row {row_index}, {name}: formula not evaluated; export values before upload."
                    )
                    value = None
                if key and value is not None:
                    mapped[key] = value
            if not mapped:
                continue
            if str(mapped.get("description", "")).strip().lower() in {
                "total",
                "grand total",
                "subtotal",
                "sub total",
            }:
                warnings.append(
                    f"{sheet} row {row_index}: summary row excluded from line items; retained in source table."
                )
                continue
            if mapped.get("invoice_number"):
                # Group within each sheet; supplier identity prevents combining different vendors.
                key = (
                    mapped["invoice_number"],
                    mapped.get("supplier_gstin") or mapped.get("supplier_name") or "",
                )
                if current_key and key[0] == current_key[0] and not key[1]:
                    key = current_key
            elif (
                current_key
                and not current_key[0].startswith("row:")
                and "invoice_number" in known
                and mapped.get("description")
                and not mapped.get("supplier_gstin")
            ):
                key = current_key
                warnings.append(
                    f"{sheet} row {row_index}: blank invoice number treated as continuation; verify grouping."
                )
            else:
                key = (f"row:{row_index}", "")
            current_key = key
            groups.setdefault(key, []).append((row_index, mapped))
        tables.append(Table(sheet=sheet, columns=names, rows=original_rows))
        for entries in groups.values():
            invoice = Invoice(
                extraction_method="tabular",
                document_type="invoice" if "invoice_number" in known else "transaction",
            )
            row_preview = ", ".join(str(i) for i, _ in entries[:100])
            if len(entries) > 100:
                row_preview += f", … ({len(entries)} rows total; last row {entries[-1][0]})"
            invoice.field_evidence = {"source": f"{sheet}, rows {row_preview}"}
            for field in (
                "invoice_number",
                "invoice_date",
                "currency",
                "place_of_supply",
                "reverse_charge",
                "document_type",
                "irn",
                "acknowledgement_number",
                "acknowledgement_date",
                "eway_bill_number",
                "vehicle_number",
                "transport_mode",
            ):
                values = list(dict.fromkeys(row[field] for _, row in entries if row.get(field)))
                if len(values) > 1:
                    warnings.append(
                        f"{sheet}: conflicting {field} for invoice {entries[0][1].get('invoice_number', '?')}: {values}"
                    )
                if values:
                    value = values[0]
                    if field == "invoice_date":
                        value = iso_date(value)
                    elif field == "reverse_charge":
                        value = {
                            "yes": True,
                            "true": True,
                            "1": True,
                            "no": False,
                            "false": False,
                            "0": False,
                        }.get(value.lower())
                        if value is None:
                            warnings.append(f"{sheet}: unrecognized reverse-charge value.")
                    elif field == "document_type":
                        value = value.lower().replace(" ", "_")
                        if value not in {"invoice", "credit_note", "debit_note", "transaction"}:
                            warnings.append(
                                f"{sheet}: unrecognized document type; retained as transaction."
                            )
                            value = "transaction"
                    if field in {
                        "irn",
                        "acknowledgement_number",
                        "acknowledgement_date",
                        "eway_bill_number",
                        "vehicle_number",
                        "transport_mode",
                    }:
                        setattr(invoice.identifiers, field, value)
                        invoice.field_evidence[f"identifiers.{field}"] = (
                            f"{sheet}, rows {row_preview}"
                        )
                    else:
                        setattr(invoice, field, value)
            for party_name in ("supplier", "buyer"):
                party = {}
                for key in ("name", "gstin", "address"):
                    values = list(
                        dict.fromkeys(
                            row[f"{party_name}_{key}"]
                            for _, row in entries
                            if row.get(f"{party_name}_{key}")
                        )
                    )
                    if values:
                        party[key] = values[0].upper() if key == "gstin" else values[0]
                        if len(values) > 1:
                            warnings.append(
                                f"{sheet}: conflicting {party_name} {key}; verify invoice grouping."
                            )
                setattr(invoice, party_name, Party(**party))
            for row_index, mapped in entries:
                item_data = {}
                item_evidence = {}
                for key in LineItem.model_fields:
                    if key not in mapped:
                        continue
                    item_evidence[key] = str(mapped[key])
                    value = (
                        mapped[key]
                        if key in {"description", "hsn_sac", "unit"}
                        else number(mapped[key])
                    )
                    if value is None:
                        warnings.append(
                            f"{sheet} row {row_index}: could not parse {key} = {mapped[key]!r}."
                        )
                    item_data[key] = value
                if item_data:
                    item = LineItem(
                        **item_data,
                        provenance={
                            "source_text": f"{sheet} row {row_index}",
                            "extraction_method": "tabular",
                            "field_evidence": item_evidence,
                        },
                    )
                    if (
                        item.taxable_value is None
                        and item.quantity is not None
                        and item.unit_price is not None
                    ):
                        item.taxable_value = item.quantity * item.unit_price - (
                            item.discount or ZERO
                        )
                        invoice.field_evidence[
                            f"line_items.{len(invoice.line_items)}.taxable_value"
                        ] = "Derived: quantity × unit_price − line discount"
                    invoice.line_items.append(item)
            for key in ("taxable_value", *TAX_KEYS, "discount"):
                values = [getattr(item, key) for item in invoice.line_items]
                if values and all(v is not None for v in values):
                    setattr(invoice.totals, key, sum(values, Decimal(0)))
                    invoice.field_evidence[f"totals.{key}"] = (
                        f"Derived: sum of {len(values)} item rows"
                    )
            for key in ("grand_total", "round_off"):
                source_values = [row[key] for _, row in entries if key in row]
                values = [number(value) for value in source_values]
                if any(value is None for value in values):
                    warnings.append(f"{sheet}: unreadable {key}; check original values.")
                values = list(dict.fromkeys(v for v in values if v is not None))
                if values:
                    setattr(invoice.totals, key, values[0])
                    if len(values) > 1:
                        warnings.append(
                            f"{sheet}: conflicting repeated invoice {key}; first value retained."
                        )
            if (
                invoice.totals.grand_total is None
                and invoice.line_items
                and all(i.total is not None for i in invoice.line_items)
            ):
                invoice.totals.grand_total = sum((i.total for i in invoice.line_items), ZERO)
                invoice.field_evidence["totals.grand_total"] = (
                    "Derived: sum of explicit line totals"
                )
            invoices.append(invoice)
            if len(invoices) > 500:
                raise ValueError("Maximum 500 invoice records per document supported.")
    if not tables:
        raise ValueError("The uploaded table is empty.")
    raw = json.dumps([table.model_dump() for table in tables], ensure_ascii=False, indent=2)
    if len(raw) > 300000:
        warnings.append("Text preview truncated; full original rows remain in the tables output.")
    return invoices, tables, raw[:300000], list(dict.fromkeys(warnings))[:200]
